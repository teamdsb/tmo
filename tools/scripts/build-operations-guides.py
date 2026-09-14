from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "docs" / "runbooks"
SALES_VISUAL_DIR = OUT_DIR / "assets" / "sales-guide"
SALES_SCREENSHOT_DIR = OUT_DIR / "assets" / "sales-screenshots"

NAVY = "17365D"
BLUE = "2E74B5"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F4F7FB"
INK = "172033"
MUTED = "5B6575"
LINE = "D9E1EA"
AMBER = "FFF4CE"
AMBER_INK = "7A5A00"
GREEN = "E7F4EA"
RED = "FDECEC"
WHITE = "FFFFFF"
PAGE_WIDTH_DXA = 9360


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], indent=120) -> None:
    total = sum(widths)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_font(run, size=None, bold=None, color=None, italic=None) -> None:
    run.font.name = "Noto Sans CJK JP"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Noto Sans CJK JP")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Noto Sans CJK JP")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK JP")
    lang = run._element.get_or_add_rPr().find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.get_or_add_rPr().append(lang)
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def set_paragraph_border(paragraph, color=LINE, size="8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, 9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    tail = paragraph.add_run(" 页")
    set_font(tail, 9, color=MUTED)


def configure_document(doc: Document, short_title: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Noto Sans CJK JP"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Noto Sans CJK JP")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Noto Sans CJK JP")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK JP")
    normal_lang = OxmlElement("w:lang")
    normal_lang.set(qn("w:val"), "zh-CN")
    normal_lang.set(qn("w:eastAsia"), "zh-CN")
    normal._element.rPr.append(normal_lang)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, before, after, color in (
        ("Heading 1", 16, 18, 10, BLUE),
        ("Heading 2", 13, 14, 7, BLUE),
        ("Heading 3", 12, 10, 5, NAVY),
    ):
        style = doc.styles[style_name]
        style.font.name = "Noto Sans CJK JP"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Noto Sans CJK JP")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Noto Sans CJK JP")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK JP")
        style_lang = OxmlElement("w:lang")
        style_lang.set(qn("w:val"), "zh-CN")
        style_lang.set(qn("w:eastAsia"), "zh-CN")
        style._element.rPr.append(style_lang)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.text = short_title
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(3)
    set_font(hp.runs[0], 8.5, bold=True, color=MUTED)
    set_paragraph_border(hp, LINE, "5")

    fp = section.footer.paragraphs[0]
    add_page_number(fp)

    props = doc.core_properties
    props.author = "TMO 项目组"
    props.company = "TMO"
    props.comments = "依据 2026-08-19 当前软件界面与工作区代码生成"


def add_cover(doc: Document, kicker: str, title: str, subtitle: str, audience: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(100)
    p.paragraph_format.space_after = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(kicker.upper())
    set_font(r, 10, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    set_font(r, 28, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(42)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    set_font(r, 13.5, color=MUTED)

    table = doc.add_table(rows=4, cols=2)
    set_table_geometry(table, [2300, 7060])
    rows = [
        ("适用对象", audience),
        ("适用版本", "当前版本（界面核对日期：2026-08-19）"),
        ("文档类型", "岗位操作手册 / 培训与日常查阅"),
        ("使用原则", "以系统实际显示和已分配权限为准"),
    ]
    for idx, (label, value) in enumerate(rows):
        for cell in table.rows[idx].cells:
            set_cell_shading(cell, PALE_BLUE if idx % 2 == 0 else WHITE)
        p0 = table.cell(idx, 0).paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        set_font(p0.add_run(label), 10.5, bold=True, color=NAVY)
        p1 = table.cell(idx, 1).paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        set_font(p1.add_run(value), 10.5, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("TMO 项目组"), 10.5, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("2026 年 8 月"), 10, color=MUTED)
    doc.add_page_break()


def add_intro(doc: Document, purpose: str, scope: list[str], not_scope: list[str]) -> None:
    doc.add_heading("使用说明", level=1)
    p = doc.add_paragraph(purpose)
    p.paragraph_format.space_after = Pt(10)
    add_callout(doc, "先看这里", "页面名称、按钮名称和角色边界均按当前代码整理。若正式环境的菜单因权限配置而隐藏，请联系管理员核对账号角色，不要共用他人账号。", LIGHT_BLUE, NAVY)
    doc.add_heading("本手册覆盖", level=2)
    add_bullets(doc, scope)
    doc.add_heading("本手册不覆盖", level=2)
    add_bullets(doc, not_scope)


def add_callout(doc: Document, label: str, text: str, fill=AMBER, color=AMBER_INK) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [PAGE_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}：")
    set_font(r, 10.5, bold=True, color=color)
    r = p.add_run(text)
    set_font(r, 10.5, color=color)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc: Document, image_path: Path, caption: str, width_inches: float = 6.45) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image_path), width=Inches(width_inches))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    cp.paragraph_format.keep_together = True
    set_font(cp.add_run(caption), 9, color=MUTED, italic=True)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        set_font(p.add_run(item), 11, color=INK)


def add_steps(doc: Document, steps: list[tuple[str, str]]) -> None:
    for idx, (title, detail) in enumerate(steps, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.38)
        p.paragraph_format.first_line_indent = Inches(-0.38)
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.keep_together = True
        set_font(p.add_run(f"{idx}. "), 11, bold=True, color=BLUE)
        set_font(p.add_run(title), 11, bold=True, color=INK)
        set_font(p.add_run(f" — {detail}"), 11, color=INK)


def add_matrix(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(header), 10, bold=True, color=NAVY)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for c_idx, value in enumerate(row):
            set_cell_shading(cells[c_idx], WHITE if r_idx % 2 == 0 else PALE_BLUE)
            p = cells[c_idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if c_idx == 0:
                set_font(p.add_run(value), 10, bold=True, color=INK)
            else:
                set_font(p.add_run(value), 10, color=INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_checklist(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.38)
        p.paragraph_format.first_line_indent = Inches(-0.38)
        p.paragraph_format.space_after = Pt(5)
        set_font(p.add_run("☐ "), 12, bold=True, color=BLUE)
        set_font(p.add_run(item), 11, color=INK)


def add_faq(doc: Document, items: list[tuple[str, str]]) -> None:
    for question, answer in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        set_font(p.add_run(f"问：{question}"), 11, bold=True, color=NAVY)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.space_after = Pt(7)
        set_font(p.add_run(f"答：{answer}"), 10.5, color=INK)


def build_sales_guide() -> Path:
    doc = Document()
    configure_document(doc, "业务员端软件功能说明")
    add_cover(doc, "Sales App Feature Guide", "业务员端软件功能说明", "功能简介与界面截图", "使用小程序业务员工作台的 SALES 账号")
    add_intro(
        doc,
        "本文只介绍当前软件页面、入口和显示内容，不对业务员的销售方法、客户管理方式或其他商业行为作要求。截图来自本地演示环境，正式环境的数据以实际账号和后台接口为准。",
        ["登录页", "业务员工作台首页", "专属推广二维码", "在线聊天与 admin 工作台", "客户列表与搜索", "订单列表", "财务页面当前状态"],
        ["销售话术、拜访流程和客户跟进方法", "价格审批、退款、发货等公司内部流程", "尚未在当前页面提供的功能"],
    )

    doc.add_heading("1. 登录", level=1)
    add_bullets(doc, [
        "登录页提供客户登录和业务员登录两个入口。",
        "使用前需要勾选隐私政策和用户服务协议。",
        "正式环境会按平台要求完成手机号授权。",
    ])
    add_figure(doc, SALES_SCREENSHOT_DIR / "01-login.png", "图 1：登录页（本地 Mock 演示环境）", 3.25)

    doc.add_heading("2. 业务员工作台首页", level=1)
    add_bullets(doc, [
        "工作台底部包含主页、客户、订单、财务四个页签。",
        "主页显示当前账号信息和专属推广二维码区域。",
        "客户扫描专属二维码并完成首次登录后，系统可建立客户归属关系。",
        "本地演示环境未分配正式业务员身份时，会显示身份或二维码提示。",
    ])
    add_figure(doc, SALES_SCREENSHOT_DIR / "02-workbench-home.png", "图 2：业务员工作台主页；截图中的红色提示来自本地演示账号。", 3.25)

    doc.add_heading("3. 在线聊天与 admin 工作台", level=1)
    add_bullets(doc, [
        "小程序在线聊天页和 admin 在线客服工作台连接的是同一套客服会话数据。",
        "客户在小程序发送的文字、图片、订单卡片或商品卡片，会显示在 admin 的同一条会话中。",
        "在 admin 中发送回复后，客户会在小程序聊天页收到消息；对客户来说，收到消息的效果与从业务员工作台发送一致。",
        "两端显示的会话状态、消息记录、客户信息和最近消息时间会同步更新。",
        "两种入口的主要区别是操作界面和权限控制，不是两套独立聊天记录。",
    ])
    add_figure(doc, SALES_SCREENSHOT_DIR / "06-miniapp-chat.png", "图 3：小程序端在线聊天页（通过电脑操作微信开发者工具取得）", 2.65)
    add_figure(doc, SALES_SCREENSHOT_DIR / "07-admin-chat.png", "图 4：admin 在线客服工作台中的同一客户会话", 6.35)
    add_matrix(doc, ["对比项", "小程序在线聊天", "admin 在线客服工作台"], [
        ["消息记录", "显示当前会话消息", "显示同一会话的完整消息记录"],
        ["消息类型", "文字、图片、订单卡片、商品卡片", "可查看并按权限发送相同类型消息"],
        ["客户信息", "以聊天页面为主", "同时显示手机号、客户 ID、归属销售和最近订单"],
        ["会话状态", "显示接待状态和当前接待角色", "显示待领取、已接单、未读和等待时间"],
        ["发送效果", "消息进入客服会话", "回复进入同一会话，客户在小程序中收到"],
        ["权限要求", "需登录并进入对应会话", "需具备后台权限；部分角色只能查看或分配，回复前可能需要认领"],
    ], [1700, 3400, 4260])
    add_callout(doc, "等同效果说明", "使用有回复权限的 admin 账号与客户聊天，和从业务员工作台的在线聊天入口发送消息，最终都进入同一客服会话；客户看到的是同一条连续聊天记录。", LIGHT_BLUE, NAVY)

    doc.add_heading("4. 客户列表", level=1)
    add_bullets(doc, [
        "客户页用于显示当前业务员权限范围内的客户。",
        "页面提供客户搜索框，可按客户信息进行查询。",
        "截图中的“客户加载失败”表示本地截图环境没有连接客户数据接口，不代表正式环境的页面结果。",
    ])
    add_figure(doc, SALES_SCREENSHOT_DIR / "03-customers.png", "图 5：客户列表与搜索入口（本地接口未连接）", 2.8)

    doc.add_heading("5. 订单列表", level=1)
    add_bullets(doc, [
        "订单页提供全部、待处理、已确认、已发货、已送达等状态筛选。",
        "订单卡片显示客户名称、订单号、日期、商品、规格、数量、金额和订单状态。",
        "页面中的具体订单内容取决于当前运行环境的数据。",
    ])
    add_figure(doc, SALES_SCREENSHOT_DIR / "04-orders.png", "图 6：订单列表（本地演示数据）", 3.25)

    doc.add_heading("6. 财务页面", level=1)
    add_bullets(doc, [
        "工作台保留财务页签。",
        "当前版本尚未接入真实佣金规则和财务结算数据。",
        "页面提示以公司正式财务系统为准。",
    ])
    add_figure(doc, SALES_SCREENSHOT_DIR / "05-finance.png", "图 7：财务页面当前显示的未接入提示", 3.25)

    doc.add_heading("7. 当前版本说明", level=1)
    add_matrix(doc, ["功能", "当前页面状态"], [
        ["业务员登录", "已提供入口；正式登录依赖平台手机号授权和账号配置"],
        ["专属推广二维码", "已提供页面区域；生成结果依赖 SALES 身份和服务端"],
        ["客户列表", "已提供列表与搜索；数据依赖客户接口和权限范围"],
        ["订单列表", "已提供筛选和订单卡片；内容依赖运行环境数据"],
        ["财务结算", "保留页签，真实结算接口尚未接入"],
        ["在线聊天", "已接入统一客服会话；小程序端与 admin 端查看和回复的是同一条会话"],
    ], [3000, 6360])
    add_callout(doc, "说明", "本文仅说明软件现状。后续功能调整后，应同步替换截图和功能描述。", LIGHT_BLUE, NAVY)

    path = OUT_DIR / "sales-miniapp-user-guide.docx"
    doc.save(path)
    return path


def build_cs_guide() -> Path:
    doc = Document()
    configure_document(doc, "客服 Admin 使用手册")
    add_cover(doc, "Customer Service Admin Guide", "客服 Admin 使用手册", "在线会话认领、回复、转接与客户上下文查询", "客服 CS；协助管理会话的 ADMIN / BOSS / MANAGER")
    add_intro(
        doc,
        "本手册用于客服坐席在 admin-web 中处理客户在线会话。标准责任角色为 CS；管理员、老板和经理可查看并协助分配，但客户沟通归口客服。",
        ["后台登录与角色确认", "消息通知与在线客服工作台", "会话筛选、优先级判断和认领", "文本、图片、订单卡片、商品卡片回复", "释放与转接", "客户资料、订单、询价和售后上下文", "异常处理和交接规范"],
        ["客户侧小程序操作", "客户归属转移、商品导入、订单派单、退款/冲正的完整流程", "权限与账号的创建配置"],
    )

    doc.add_heading("1. 登录与权限", level=1)
    add_steps(doc, [
        ("打开后台登录页", "输入管理员分配的用户名和密码。"),
        ("处理本地会话提示", "共享设备上看到“检测到本地登录会话”时，核对账号与角色；不是本人则点击“清除并重新登录”。"),
        ("选择角色", "一个账号存在多个后台角色时，在弹窗中选择“客服（CS）”。后台支持 BOSS、ADMIN、CS、MANAGER。"),
        ("确认进入后台", "在侧栏底部核对用户名和角色，客服日常操作应使用 CS。"),
    ])
    add_callout(doc, "权限边界", "CS 负责售后、询单、发运与物流相关工作，订单在客服角色下以读取为主；确认线下收款、派单和发货前改派不属于 CS 权限。", AMBER, AMBER_INK)

    doc.add_heading("2. 进入在线客服工作台", level=1)
    add_bullets(doc, [
        "优先从顶部消息铃铛查看待处理消息；点击通知项会直接打开对应会话。",
        "也可从部署提供的“在线客服工作台”入口进入实时会话页；页面标题为“在线客服工作台”。",
        "若侧栏“在线客服”进入的是询价列表，应通过顶部通知进入实时会话工作台，或联系管理员确认当前部署入口。",
        "进入后可点击右上角“刷新”重新加载会话、员工和商品参考数据。",
    ])
    add_callout(doc, "环境提示", "演示/Mock 环境中的认领、释放和转接按钮会被禁用。正式联调或生产操作必须使用真实环境；演示数据不能作为客户服务记录。", LIGHT_BLUE, NAVY)

    doc.add_heading("3. 工作台页面地图", level=1)
    add_matrix(doc, ["区域", "内容", "使用要点"], [
        ["左栏：会话列表", "全部 / 我的会话 / 待领取 / 未读", "按等待、未读和最近消息选择任务"],
        ["中栏：聊天区", "消息记录、认领/释放、回复工具", "完成实际沟通"],
        ["右栏：客户资料", "姓名、手机号、客户 ID、归属销售、当前坐席、状态", "回复前核对客户身份"],
        ["右栏：业务上下文", "最近订单、询价、售后", "减少重复询问并核对问题对象"],
        ["顶部：通知", "待处理数量、消息预览、超时提示", "快速跳转到目标会话"],
    ], [1900, 3400, 4060])

    doc.add_heading("4. 会话筛选与优先级", level=1)
    add_matrix(doc, ["筛选项", "显示范围", "推荐用法"], [
        ["全部", "可访问的全部会话", "管理或查找特定客户"],
        ["我的会话", "当前坐席已认领的会话", "处理自己的在办任务"],
        ["待领取", "尚未分配坐席的会话", "CS 默认入口，优先认领"],
        ["未读", "有客服侧未读消息的会话", "避免漏回客户"],
    ], [1700, 3500, 4160])
    add_bullets(doc, [
        "待领取会话会显示“已等待”计时；等待达到 30 秒后显示“等待超时”，并优先排在列表上方。",
        "“待回复 N”表示客服侧未读数量。打开会话后系统会尝试自动标记已读。",
        "同等优先级下，按最近消息时间处理；涉及付款、发货、投诉升级时结合公司服务规范提级。",
    ])

    doc.add_heading("5. 标准处理流程", level=1)
    add_steps(doc, [
        ("选择会话", "优先从“待领取”或“未读”选择等待超时、待回复数量高的会话。"),
        ("核对身份与上下文", "查看右栏姓名、手机号、客户 ID、归属销售、最近订单、询价和售后。"),
        ("点击“认领”", "成功后会话状态变为“已接单”，当前坐席更新为本人。"),
        ("阅读完整消息", "区分客户消息、客服消息和系统消息；确认客户具体诉求。"),
        ("发送回复", "输入明确答复，或发送图片、订单卡片、商品卡片辅助说明。"),
        ("确认发送结果", "消息应出现在聊天区；失败时不要连续重复点击，先读取页面提示。"),
        ("继续跟进或交接", "本人继续处理则保留在“我的会话”；需他人接手时转接，暂不处理时按规范释放。"),
    ])

    doc.add_heading("6. 发送不同类型的消息", level=1)
    add_matrix(doc, ["类型", "操作", "注意事项"], [
        ["文本", "在“输入回复内容...”输入后点击发送图标", "空白内容不能发送；先核对客户和订单"],
        ["图片", "点击“发送图片”并选择文件", "仅 jpg/png/webp，最大 5MB；不得上传敏感或无关图片"],
        ["订单卡片", "点击最近订单对应的“发订单卡片”", "页面最多提供最近 2 条快捷按钮；发送前核对订单号和状态"],
        ["商品卡片", "点击“发商品卡片”", "页面最多提供前 2 个商品快捷按钮；客户可查看商品，客服端卡片可进入商品详情"],
        ["系统消息", "系统自动产生", "用于识别认领、释放或转接等状态变化，不作为人工回复"],
    ], [1500, 3600, 4260])
    add_callout(doc, "发送规范", "不要只回复“已处理”而不给结果或下一步。推荐结构：确认问题 → 给出结论/动作 → 说明预计时间 → 必要时附订单或商品卡片。", GREEN, "245A32")

    doc.add_heading("7. 认领、释放与转接", level=1)
    add_matrix(doc, ["动作", "适用场景", "结果与限制"], [
        ["认领", "待领取会话由本人接手", "状态从“待领取”变为“已接单”；正式环境可用"],
        ["释放", "会话需要回到公共队列", "清除当前坐席并重新计入等待；避免无说明释放"],
        ["转接", "需要另一名客服继续处理", "先选目标坐席再点“转移会话”；目标必须是已启用的 CS 账号"],
    ], [1500, 3400, 4460])
    add_bullets(doc, [
        "CS、ADMIN、BOSS、MANAGER 可执行会话转接；客户沟通仍归口 CS。",
        "释放或转接前应在聊天中给客户必要说明，并在团队交接渠道补充背景。",
        "不要把会话转给 SALES：业务员只作为客户归属背景展示，不进入 admin 客服工作台。",
        "已关闭会话在当前页面没有关闭按钮；不要用释放替代正式关闭流程。",
    ])

    doc.add_heading("8. 使用客户上下文", level=1)
    add_matrix(doc, ["信息", "用途", "核对重点"], [
        ["客户资料", "确认正在服务的人", "姓名、手机号、客户 ID"],
        ["归属销售", "了解客户负责人", "只作背景，不把会话转给业务员"],
        ["最近订单", "处理催单、物流、订单问题", "订单号、首件商品、状态"],
        ["询价", "了解售前价格沟通", "询价编号、状态、客户留言"],
        ["售后", "了解已有售后问题", "工单编号、状态、主题"],
    ], [1700, 3300, 4360])
    add_callout(doc, "隐私要求", "客户手机号、客户 ID、订单和售后内容仅用于本次工作。不得复制到无关群聊或个人工具；截图报障时遮挡与问题无关的信息。", AMBER, AMBER_INK)

    doc.add_heading("9. 通知与已读机制", level=1)
    add_bullets(doc, [
        "顶部铃铛显示待处理数量，超过 99 时显示“99+”。",
        "新消息会出现短暂提示；点击提示或通知项可直接打开对应会话。",
        "待领取会话等待超过 30 秒会产生超时提示，应优先处理。",
        "打开含未读消息的会话后系统会自动标记已读；若数字未清除，点击“刷新”并确认网络。",
        "当前 admin 浏览器环境以轮询刷新为兜底，因此新消息可能有短暂延迟。",
    ])

    doc.add_heading("10. 推荐话术结构", level=1)
    add_matrix(doc, ["场景", "推荐表达"], [
        ["首次响应", "您好，我是客服××，已看到您关于【订单/商品】的问题，我先为您核对。"],
        ["需要等待", "已提交核对，预计在【时间】前回复您；如有更新我会在本会话同步。"],
        ["发送订单卡片", "请确认下方订单卡片是否为您咨询的订单，我将按该订单继续处理。"],
        ["转接", "该问题需要由另一位客服继续处理，我已转接并同步背景，请稍候。"],
        ["暂无法确认", "目前信息不足，我需要核对【具体事项】；在确认前不建议您进行【风险操作】。"],
    ], [1900, 7460])

    doc.add_heading("11. 交接与收班检查", level=1)
    add_checklist(doc, [
        "“我的会话”中没有无人跟进的紧急消息。",
        "等待超时会话已认领或明确转接。",
        "已承诺回复时间的会话已记录并交接。",
        "需要其他客服处理的会话已转给已启用的 CS 账号，不只是口头通知。",
        "发送失败、图片失败或数据异常已记录账号、会话 ID、时间和提示。",
        "共享设备已退出登录，浏览器未保留他人可继续使用的会话。",
    ])

    doc.add_heading("12. 常见问题与处理", level=1)
    add_faq(doc, [
        ("登录提示账号或密码错误？", "重新输入并确认大小写；仍失败则联系 BOSS 或账号管理员重置，不要借用他人账号。"),
        ("看不到在线客服工作台或通知？", "确认当前角色为 CS 且拥有对应权限；尝试刷新页面。若侧栏进入询价页，可通过顶部通知进入实时会话。"),
        ("认领、释放、转接按钮是灰色？", "可能未选择会话、当前为 Mock 环境、没有目标坐席，或权限/角色不符合。先核对页面和环境。"),
        ("图片上传失败？", "确认格式为 jpg/png/webp 且不超过 5MB；压缩后重试。不要连续上传同一文件。"),
        ("消息发送失败？", "保留输入内容，检查网络和页面提示，刷新会话确认是否实际发送成功后再重试，避免重复消息。"),
        ("转接列表没有目标客服？", "只有状态为“已启用”且角色包含 CS 的员工会出现。联系管理员核对对方账号状态与角色。"),
        ("会话显示不存在或无法访问？", "可能链接过期、会话已不可见或权限不足。返回会话列表重新查找，并记录会话 ID 报障。"),
        ("未读数没有清零？", "打开会话等待自动标记已读，再点击刷新；持续异常时记录会话 ID 和未读数。"),
    ])

    doc.add_heading("附录 A：状态与标识", level=1)
    add_matrix(doc, ["页面标识", "系统含义", "客服动作"], [
        ["待领取", "OPEN_UNASSIGNED，尚无坐席", "优先认领"],
        ["已接单", "OPEN_ASSIGNED，已有坐席", "由当前坐席跟进或转接"],
        ["已关闭", "CLOSED", "只读核对；当前工作台无关闭按钮"],
        ["待回复 N", "客服侧未读消息数", "打开并回复"],
        ["已等待", "待领取队列等待时长", "关注时效"],
        ["等待超时", "等待达到 30 秒", "立即优先处理"],
    ], [2200, 3800, 3360])

    doc.add_heading("附录 B：角色责任速查", level=1)
    add_matrix(doc, ["角色", "在客服工作台中的责任"], [
        ["CS", "标准客服坐席：认领、回复、释放、转接，处理客户沟通"],
        ["ADMIN / BOSS / MANAGER", "可查看并协助管理或分配；沟通归口 CS"],
        ["SALES", "仅显示为客户归属背景，不进入 admin 客服工作台"],
        ["CUSTOMER", "从小程序发起并查看自己的客服会话"],
    ], [2800, 6560])

    path = OUT_DIR / "customer-service-admin-user-guide.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for output in (build_sales_guide(), build_cs_guide()):
        print(output)
